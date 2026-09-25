"""Optional real LibreOffice integration QA; outputs are synthetic fixtures in .build/."""
from pathlib import Path
import sys
import threading
import urllib.request

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / '.build/qa-python')]
import server
import fitz
from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw


def main():
    out = ROOT / '.build/qa-office'
    out.mkdir(parents=True, exist_ok=True)
    picture = Image.new('RGB', (400, 100), '#214d40')
    ImageDraw.Draw(picture).text((20, 40), 'MyFlipbook - layout test', fill='white')
    picture.save(out / 'banner.png')
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = 'SARVAMAYA - document header'
    doc.add_heading('Word layout verification', 0)
    doc.add_paragraph('This document includes a real table, an image and an explicit page break.')
    doc.add_picture(str(out / 'banner.png'), width=Inches(4))
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    for cell, value in zip(table.rows[0].cells, ['Item', 'Quantity', 'Total']):
        cell.text = value
    for data in [('Books', '3', '150'), ('Folders', '2', '40')]:
        for cell, value in zip(table.add_row().cells, data):
            cell.text = value
    doc.add_page_break()
    doc.add_heading('Second page', 1)
    doc.add_paragraph('END_WORD_LAYOUT')
    doc.save(out / 'layout.docx')
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Sales'
    sheet.append(['Item', 'Quantity', 'Price'])
    for i in range(1, 76):
        sheet.append([f'ROW_{i}', i, i * 10])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='214D40')
    for col in 'ABC':
        sheet.column_dimensions[col].width = 22
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = '1:1'
    sheet.print_area = 'A1:C76'
    second = wb.create_sheet('Summary')
    second.append(['SECOND_SHEET', 'Formula'])
    second.append(['END_EXCEL_LAYOUT', '=SUM(Sales!B2:B76)'])
    second.column_dimensions['A'].width = 30
    second.column_dimensions['B'].width = 20
    wb.save(out / 'layout.xlsx')
    (out / 'data.csv').write_text('Name;Value\nCaf\u00e9;42\nCSV_END;99\n', encoding='utf-8')
    httpd = server.ThreadingHTTPServer(('127.0.0.1', 0), server.Handler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    try:
        for filename, route, expected in [
            ('layout.docx', 'word', ['END_WORD_LAYOUT', 'Quantity']),
            ('layout.xlsx', 'excel', ['ROW_75', 'SECOND_SHEET', '2850']),
            ('data.csv', 'excel', ['CSV_END', '42']),
        ]:
            source = out / filename
            mime = next(k for k, v in server.OFFICE_MIMES.items() if v == source.suffix)
            request = urllib.request.Request(f'http://127.0.0.1:{httpd.server_port}/api/convert/{route}-to-pdf',
                data=source.read_bytes(), headers={'Content-Type':mime, 'X-Filename':filename, 'X-Build-Token':server.TOKEN})
            with urllib.request.urlopen(request, timeout=300) as response:
                payload = response.read()
            destination = out / (filename + '.pdf')
            destination.write_bytes(payload)
            pdf = fitz.open(destination)
            text = '\n'.join(page.get_text() for page in pdf)
            for marker in expected:
                assert marker in text, (filename, marker, text)
            if source.suffix == '.docx':
                assert len(pdf) == 2
                assert pdf[0].get_images(), 'Word image was lost'
            if source.suffix == '.xlsx':
                assert pdf[0].rect.width > pdf[0].rect.height, 'Landscape print setting lost'
            for i, page in enumerate(pdf):
                page.get_pixmap(matrix=fitz.Matrix(1, 1)).save(out / f'{filename}-{i+1}.png')
            print(f'PASS real HTTP + LibreOffice: {filename}, {len(pdf)} PDF pages, complete markers')
    finally:
        httpd.shutdown()
        httpd.server_close()
        thread.join()


if __name__ == '__main__':
    main()
